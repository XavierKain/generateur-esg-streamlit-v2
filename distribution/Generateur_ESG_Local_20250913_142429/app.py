#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GÉNÉRATEUR ESG SIMPLIFIÉ - Interface Streamlit
Application simplifiée pour la génération de questionnaires ESG uniquement
"""

import streamlit as st
import os
import tempfile
import shutil
import datetime
import time
import glob
import zipfile
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
import warnings
from xlwings_generator import XLWingsGenerator

# Configuration de la page
st.set_page_config(
    page_title="Générateur ESG - Simple",
    page_icon="🏭",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Suppression des warnings
warnings.filterwarnings('ignore')

# CSS personnalisé pour améliorer l'apparence avec effets hover
st.markdown("""
<style>
    .main-header {
        font-size: 2.5rem;
        color: #1f77b4;
        text-align: center;
        margin-bottom: 2rem;
        padding: 1rem;
        background: linear-gradient(90deg, #f0f8ff, #e6f3ff);
        border-radius: 10px;
    }
    .step-header {
        font-size: 1.5rem;
        color: #2e8b57;
        margin: 1rem 0;
        padding: 0.5rem;
        background-color: #f0fff0;
        border-left: 4px solid #2e8b57;
        border-radius: 5px;
    }
    .success-box {
        padding: 1rem;
        background-color: #d4edda;
        border: 1px solid #c3e6cb;
        border-radius: 5px;
        color: #155724;
        margin: 1rem 0;
    }
    .warning-box {
        padding: 1rem;
        background-color: #fff3cd;
        border: 1px solid #ffeaa7;
        border-radius: 5px;
        color: #856404;
        margin: 1rem 0;
    }
    .error-box {
        padding: 1rem;
        background-color: #f8d7da;
        border: 1px solid #f5c6cb;
        border-radius: 5px;
        color: #721c24;
        margin: 1rem 0;
    }
    /* Effets hover pour tous les boutons */
    .stButton > button:hover {
        transform: translateY(-2px);
        box-shadow: 0 4px 8px rgba(0,0,0,0.2);
        transition: all 0.3s ease;
    }
    .stButton > button {
        transition: all 0.3s ease;
    }
    /* Barre de progression pleine largeur */
    .stProgress > div > div > div > div {
        width: 100% !important;
    }
</style>
""", unsafe_allow_html=True)

def init_session_state():
    """Initialiser les variables de session"""
    if 'temp_dir' not in st.session_state:
        st.session_state.temp_dir = tempfile.mkdtemp()
    if 'uploaded_bdd_file' not in st.session_state:
        st.session_state.uploaded_bdd_file = None
    if 'uploaded_template_file' not in st.session_state:
        st.session_state.uploaded_template_file = None
    if 'selected_year' not in st.session_state:
        st.session_state.selected_year = None
    if 'available_years' not in st.session_state:
        st.session_state.available_years = []
    if 'generated_questionnaires' not in st.session_state:
        st.session_state.generated_questionnaires = False
    if 'current_step' not in st.session_state:
        st.session_state.current_step = 1
    if 'generation_result' not in st.session_state:
        st.session_state.generation_result = None
    if 'generated_zip_data' not in st.session_state:
        st.session_state.generated_zip_data = None
    if 'selected_questionnaire_indices' not in st.session_state:
        st.session_state.selected_questionnaire_indices = []
    if 'questionnaires_preview' not in st.session_state:
        st.session_state.questionnaires_preview = []

def main():
    init_session_state()
    
    # Titre principal
    st.markdown("""
    <div class="main-header">
        🏭 GÉNÉRATEUR ESG SIMPLIFIÉ
        <br><small>Génération automatique de questionnaires ESG</small>
    </div>
    """, unsafe_allow_html=True)
    
    # Interface de génération
    page_generation_workflow()

def page_generation_workflow():
    """Interface principale de génération"""
    
    # Navigation par étapes
    st.markdown('<div class="step-header">🏭 GÉNÉRATION DE QUESTIONNAIRES</div>', unsafe_allow_html=True)
    
    # Interface progressive par étapes
    if st.session_state.current_step == 1:
        page_upload_files()
    elif st.session_state.current_step == 2:
        page_generation_config()
    elif st.session_state.current_step == 3:
        page_generation_execute()
    
    # Barre de progression avec navigation
    st.markdown("---")
    progress_indicators = st.columns(3)
    with progress_indicators[0]:
        if st.session_state.current_step >= 1:
            if st.session_state.current_step == 1:
                st.warning("🔄 1. Upload Fichiers")
            else:
                if st.button("✅ 1. Upload Fichiers", help="Cliquer pour revenir à cette étape"):
                    st.session_state.current_step = 1
                    st.rerun()
        else:
            st.info("⏸️ 1. Upload Fichiers")
    
    with progress_indicators[1]:
        if st.session_state.current_step >= 2:
            if st.session_state.current_step == 2:
                st.warning("🔄 2. Configuration")
            else:
                if st.button("✅ 2. Configuration", help="Cliquer pour revenir à cette étape"):
                    st.session_state.current_step = 2
                    st.rerun()
        else:
            st.info("⏸️ 2. Configuration")
    
    with progress_indicators[2]:
        if st.session_state.current_step >= 3:
            st.warning("🔄 3. Génération")
        else:
            st.info("⏸️ 3. Génération")
    
    # Bouton de remise à zéro
    if st.session_state.current_step > 1:
        st.markdown("---")
        col1, col2, col3 = st.columns([1, 1, 2])
        with col1:
            if st.button("🔄 Recommencer", type="secondary", help="Remettre à zéro et recommencer depuis le début"):
                st.session_state.current_step = 1
                st.session_state.uploaded_bdd_file = None
                st.session_state.uploaded_template_file = None
                st.session_state.selected_year = "2025"
                st.session_state.generated_questionnaires = False
                st.session_state.generated_zip_data = None
                st.session_state.selected_questionnaire_indices = []
                st.session_state.questionnaires_preview = []
                st.success("✅ Application remise à zéro")
                st.rerun()

def page_upload_files():
    """Page upload des fichiers"""
    st.markdown("**📁 Upload des fichiers requis**")
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.markdown("**1. Fichier BDD ESG**")
        uploaded_bdd = st.file_uploader("Sélectionnez votre fichier BDD", type=['xlsx', 'xlsm'], key="bdd_upload")
        if uploaded_bdd:
            # Sauvegarder dans le répertoire temporaire
            bdd_path = os.path.join(st.session_state.temp_dir, uploaded_bdd.name)
            with open(bdd_path, "wb") as f:
                f.write(uploaded_bdd.getbuffer())
            st.session_state.uploaded_bdd_file = bdd_path
            st.success("✅ Fichier BDD uploadé")
    
    with col2:
        st.markdown("**2. Template questionnaire**")
        uploaded_template = st.file_uploader("Sélectionnez votre template", type=['xlsx', 'xlsm'], key="template_upload")
        if uploaded_template:
            # Sauvegarder dans le répertoire temporaire
            template_path = os.path.join(st.session_state.temp_dir, uploaded_template.name)
            with open(template_path, "wb") as f:
                f.write(uploaded_template.getbuffer())
            st.session_state.uploaded_template_file = template_path
            st.success("✅ Template uploadé")
    
    # Analyser le fichier BDD si les deux sont uploadés
    if st.session_state.uploaded_bdd_file and st.session_state.uploaded_template_file:
        try:
            wb = load_workbook(st.session_state.uploaded_bdd_file, data_only=True)
            sheet_names = wb.sheetnames
            wb.close()
            
            # Détecter automatiquement les années (onglets qui sont des nombres de 4 chiffres)
            available_years = []
            for sheet in sheet_names:
                if sheet.isdigit() and len(sheet) == 4 and 2020 <= int(sheet) <= 2030:
                    available_years.append(sheet)
            
            available_years.sort()
            
            if not available_years:
                st.error("❌ Aucun onglet d'année détecté (format attendu: 2023, 2024, etc.)")
                return
            else:
                st.success(f"✅ Fichiers validés - Années détectées: {', '.join(available_years)}")
                st.session_state.available_years = available_years
                # Sélectionner la dernière année par défaut
                st.session_state.selected_year = available_years[-1]
                
                # Bouton pour passer à l'étape suivante
                if st.button("➡️ Passer à la configuration", type="primary"):
                    st.session_state.current_step = 2
                    st.rerun()
                    
        except Exception as e:
            st.error(f"❌ Erreur lors de l'analyse: {e}")

def page_generation_config():
    """Page configuration pour génération"""
    if not st.session_state.uploaded_bdd_file or not st.session_state.uploaded_template_file:
        st.warning("⚠️ Veuillez d'abord uploader les fichiers")
        return
    
    if not st.session_state.available_years:
        st.error("❌ Aucune année disponible détectée")
        return
    
    st.markdown("**⚙️ Configuration de la génération**")
    
    # Sélection de l'année avec radio buttons basé sur les années détectées
    st.markdown("**📅 Sélection de l'année :**")
    
    # Assurer que selected_year est dans la liste des années disponibles
    if st.session_state.selected_year not in st.session_state.available_years:
        st.session_state.selected_year = st.session_state.available_years[-1]
    
    selected_year = st.radio(
        "Choisissez l'année",
        options=st.session_state.available_years,
        index=st.session_state.available_years.index(st.session_state.selected_year),
        horizontal=True,
        key="year_selection"
    )
    
    # Mettre à jour l'année sélectionnée
    if selected_year != st.session_state.selected_year:
        st.session_state.selected_year = selected_year
        st.rerun()
    
    st.write(f"**Année sélectionnée :** {st.session_state.selected_year}")
    
    # Status xlwings - section informative
    st.markdown("**🔧 Status du système de formatage :**")
    with st.expander("ℹ️ Informations sur les moteurs de formatage", expanded=False):
        try:
            xlwings_gen = XLWingsGenerator()
            xlwings_available, xlwings_msg = xlwings_gen.is_available()
            
            col1, col2 = st.columns(2)
            with col1:
                st.markdown("**🎨 xlwings (Recommandé)**")
                if xlwings_available:
                    st.success(f"✅ {xlwings_msg}")
                    st.info("🎯 **Avantage** : Conserve 100% du formatage conditionnel")
                else:
                    st.error(f"❌ {xlwings_msg}")
                    
            with col2:
                st.markdown("**🔧 openpyxl (Fallback)**")
                st.warning("⚠️ Formatage conditionnel perdu")
                st.info("📝 Utilisé automatiquement si xlwings indisponible")
                
            if xlwings_available:
                st.success("🚀 **xlwings détecté** - Formatage conditionnel préservé !")
            else:
                st.warning("⚠️ **xlwings indisponible** - Utilisation d'openpyxl (formatage basique)")
                
        except Exception as e:
            st.error(f"❌ Erreur lors de la vérification xlwings: {e}")
    
    # Prévisualisation des données et sélection des questionnaires
    with st.expander("📊 Prévisualiser et sélectionner les questionnaires à générer", expanded=True):
        try:
            # Import de la nouvelle fonction
            from generateur_2025_streamlit import get_questionnaires_preview
            
            questionnaires_preview = get_questionnaires_preview(st.session_state.uploaded_bdd_file, st.session_state.selected_year)
            st.session_state.questionnaires_preview = questionnaires_preview
            
            if questionnaires_preview:
                st.write(f"**{len(questionnaires_preview)} questionnaires disponibles pour {st.session_state.selected_year}**")
                
                # Boutons de sélection rapide
                col1, col2, col3 = st.columns(3)
                with col1:
                    if st.button("✅ Tout sélectionner"):
                        st.session_state.selected_questionnaire_indices = list(range(len(questionnaires_preview)))
                        st.rerun()
                with col2:
                    if st.button("❌ Tout désélectionner"):
                        st.session_state.selected_questionnaire_indices = []
                        st.rerun()
                with col3:
                    if st.button("🔄 Inverser la sélection"):
                        all_indices = set(range(len(questionnaires_preview)))
                        current_selection = set(st.session_state.selected_questionnaire_indices)
                        st.session_state.selected_questionnaire_indices = list(all_indices - current_selection)
                        st.rerun()
                
                # Affichage des questionnaires avec checkboxes
                st.markdown("**📋 Sélectionnez les questionnaires à générer :**")
                
                # Afficher par lots pour éviter une interface trop lourde
                items_per_page = 20
                total_pages = (len(questionnaires_preview) + items_per_page - 1) // items_per_page
                
                if total_pages > 1:
                    current_page = st.selectbox("📄 Page", range(1, total_pages + 1), index=0) - 1
                else:
                    current_page = 0
                
                start_idx = current_page * items_per_page
                end_idx = min(start_idx + items_per_page, len(questionnaires_preview))
                
                for i in range(start_idx, end_idx):
                    questionnaire = questionnaires_preview[i]
                    
                    # Checkbox pour chaque questionnaire
                    is_selected = i in st.session_state.selected_questionnaire_indices
                    
                    col1, col2 = st.columns([1, 4])
                    with col1:
                        selected = st.checkbox("Sélect.", value=is_selected, key=f"quest_{i}", label_visibility="hidden")
                    
                    with col2:
                        # Affichage des informations
                        st.markdown(f"""
                        **#{questionnaire['id']} - {questionnaire['locataire']}**  
                        📍 {questionnaire['adresse']}  
                        📅 {questionnaire['date']}  
                        📁 `{questionnaire['filename_preview']}`
                        """)
                    
                    # Mettre à jour la sélection
                    if selected and i not in st.session_state.selected_questionnaire_indices:
                        st.session_state.selected_questionnaire_indices.append(i)
                    elif not selected and i in st.session_state.selected_questionnaire_indices:
                        st.session_state.selected_questionnaire_indices.remove(i)
                
                # Résumé de la sélection
                selected_count = len(st.session_state.selected_questionnaire_indices)
                if selected_count > 0:
                    st.success(f"✅ {selected_count} questionnaire(s) sélectionné(s) sur {len(questionnaires_preview)}")
                else:
                    st.warning("⚠️ Aucun questionnaire sélectionné")
                    
            else:
                st.warning(f"Aucune donnée trouvée pour l'année {st.session_state.selected_year}")
        except Exception as e:
            st.error(f"Erreur lors de la prévisualisation: {e}")
    
    # Ancienne prévisualisation simple (optionnelle)
    with st.expander("📊 Prévisualisation simple (premiers éléments)", expanded=False):
        try:
            df = preview_data(st.session_state.uploaded_bdd_file, st.session_state.selected_year)
            if df is not None and not df.empty:
                st.write(f"**{len(df)} propriétés trouvées pour {st.session_state.selected_year}**")
                st.dataframe(df.head(10), use_container_width=True)
            else:
                st.warning(f"Aucune donnée trouvée pour l'année {st.session_state.selected_year}")
        except Exception as e:
            st.error(f"Erreur lors de la prévisualisation: {e}")
    
    # Bouton pour passer à l'étape suivante
    st.markdown("---")
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        selected_count = len(st.session_state.selected_questionnaire_indices)
        if selected_count > 0:
            if st.button(f"✅ Valider et générer {selected_count} questionnaire(s)", type="primary", use_container_width=True):
                st.success(f"✅ Configuration validée ! {selected_count} questionnaire(s) sélectionné(s)")
                # Passer automatiquement à l'étape suivante
                st.session_state.current_step = 3
                st.info("➡️ Passage automatique à l'étape Génération...")
                time.sleep(1.5)
                st.rerun()
        else:
            st.warning("⚠️ Veuillez sélectionner au moins un questionnaire pour continuer")

def page_generation_execute():
    """Page exécution génération"""
    if not st.session_state.uploaded_bdd_file or not st.session_state.uploaded_template_file:
        st.warning("⚠️ Veuillez d'abord uploader les fichiers")
        return
    
    st.markdown("**🚀 Génération des questionnaires**")
    
    # Résumé configuration
    selected_count = len(st.session_state.selected_questionnaire_indices)
    st.markdown(f"""
    **Configuration actuelle :**
    - 📁 Fichier BDD: `{os.path.basename(st.session_state.uploaded_bdd_file)}`
    - 📋 Template: `{os.path.basename(st.session_state.uploaded_template_file)}`
    - 📅 Année: `{st.session_state.selected_year}`
    - 📊 Questionnaires sélectionnés: `{selected_count}`
    """)
    
    # Afficher les résultats si la génération a déjà été faite
    if st.session_state.generated_questionnaires and st.session_state.generation_result:
        display_generation_results()
    else:
        # Bouton de génération avec largeur normale et alignement à gauche
        col1, col2, col3 = st.columns([1, 1, 2])
        with col1:
            selected_count = len(st.session_state.selected_questionnaire_indices)
            if selected_count > 0:
                if st.button(f"🚀 GÉNÉRER {selected_count} QUESTIONNAIRE(S)", type="primary"):
                    generate_questionnaires()
            else:
                st.error("❌ Aucun questionnaire sélectionné")

def display_generation_results():
    """Afficher les résultats de génération persistants"""
    result = st.session_state.generation_result
    
    st.success("✅ Génération terminée avec succès !")
    
    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Questionnaires générés", result['generated_count'])
    with col2:
        st.metric("Échecs", result['failed_count'])
    with col3:
        if 'total_available' in result:
            st.metric("Total disponible", result['total_available'])
    
    # Information sur la sélection
    if 'total_available' in result and result['total_available'] > result['total_processed']:
        st.info(f"ℹ️ {result['total_processed']} questionnaire(s) sélectionné(s) sur {result['total_available']} disponibles")
    
    # Bouton de téléchargement ZIP avec données directement intégrées
    if st.session_state.generated_zip_data:
        timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        selected_count = result.get('total_processed', 'selection')
        zip_filename = f"questionnaires_{st.session_state.selected_year}_{selected_count}items_{timestamp}.zip"
        
        st.download_button(
            label=f"📥 Télécharger le ZIP des questionnaires ({result['generated_count']} fichiers)",
            data=st.session_state.generated_zip_data,
            file_name=zip_filename,
            mime="application/zip",
            type="secondary",
            help=f"Télécharger les {result['generated_count']} questionnaires générés pour l'année {st.session_state.selected_year}"
        )
        
    # Bouton pour recommencer après génération
    st.markdown("---")
    col1, col2, col3 = st.columns([1, 1, 2])
    with col1:
        if st.button("🔄 Nouvelle génération", type="secondary", help="Recommencer une nouvelle génération"):
            st.session_state.current_step = 1
            st.session_state.uploaded_bdd_file = None
            st.session_state.uploaded_template_file = None
            st.session_state.selected_year = "2025"
            st.session_state.generated_questionnaires = False
            st.session_state.generation_result = None
            st.session_state.generated_zip_data = None
            st.session_state.selected_questionnaire_indices = []
            st.session_state.questionnaires_preview = []
            st.success("✅ Prêt pour une nouvelle génération")
            st.rerun()

def create_and_show_download_button(output_dir):
    """Créer et afficher le bouton de téléchargement avec données intégrées"""
    try:
        # Nom unique pour le ZIP basé sur le timestamp et l'année
        timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        zip_filename = f"questionnaires_{st.session_state.selected_year}_{timestamp}.zip"
        zip_path = os.path.join(st.session_state.temp_dir, zip_filename)
        
        # Créer le ZIP avec uniquement les fichiers de cette génération
        if not os.path.exists(zip_path):
            with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                # Vérifier que le dossier de sortie existe et contient des fichiers
                if os.path.exists(output_dir):
                    files_added = 0
                    for root, dirs, files in os.walk(output_dir):
                        for file in files:
                            if file.endswith(('.xlsx', '.xlsm', '.pdf', '.txt')):  # Filtrer les types de fichiers
                                file_path = os.path.join(root, file)
                                # Nom dans le ZIP relatif au dossier de génération
                                arcname = os.path.relpath(file_path, output_dir)
                                zipf.write(file_path, arcname)
                                files_added += 1
                    
                    if files_added == 0:
                        st.warning("⚠️ Aucun fichier questionnaire trouvé dans le dossier de génération")
                        return
                    else:
                        st.info(f"📦 ZIP créé avec {files_added} fichiers")
                else:
                    st.error("❌ Dossier de génération introuvable")
                    return
        
        # Lire les données du ZIP
        with open(zip_path, "rb") as f:
            zip_data = f.read()
        
        # Bouton de téléchargement direct
        st.download_button(
            label=f"📥 Télécharger le ZIP des questionnaires ({st.session_state.selected_year})",
            data=zip_data,
            file_name=zip_filename,
            mime="application/zip",
            type="secondary",
            help=f"Télécharger uniquement les questionnaires générés pour l'année {st.session_state.selected_year}"
        )
            
    except Exception as e:
        st.error(f"❌ Erreur lors de la création du ZIP: {e}")

def preview_data(bdd_file, year):
    """Prévisualiser les données d'une année donnée"""
    try:
        wb = load_workbook(bdd_file, data_only=True)
        if year not in wb.sheetnames:
            return None
            
        ws = wb[year]
        data = []
        
        for row in range(10, min(50, ws.max_row + 1)):  # Prévisualisation limitée
            # Colonne O = nom du dossier (colonne 15)
            folder_name = ws.cell(row=row, column=15).value
            if not folder_name or str(folder_name).strip() == '':
                continue
                
            # Autres colonnes importantes
            numero_identification = ws.cell(row=row, column=1).value
            locataire = ws.cell(row=row, column=6).value
            adresse = ws.cell(row=row, column=4).value
            
            data.append({
                'N° ID': numero_identification,
                'Dossier': str(folder_name).strip(),
                'Locataire': str(locataire).strip() if locataire else '',
                'Adresse': str(adresse).strip() if adresse else ''
            })
        
        wb.close()
        return pd.DataFrame(data)
        
    except Exception as e:
        st.error(f"Erreur lors de la lecture des données: {e}")
        return None

def generate_questionnaires():
    """Générer les questionnaires avec barre de progression et création directe du ZIP"""
    
    # Barre de progression pleine largeur
    progress_bar = st.progress(0)
    status_text = st.empty()
    
    try:
        # Import du module de génération
        from generateur_2025_streamlit import generate_selected_questionnaires_to_zip
        
        # Paramètres
        bdd_file = st.session_state.uploaded_bdd_file
        template_file = st.session_state.uploaded_template_file
        year = st.session_state.selected_year
        selected_indices = st.session_state.selected_questionnaire_indices
        
        if not selected_indices:
            st.error("❌ Aucun questionnaire sélectionné")
            return
        
        st.info(f"📁 Génération de {len(selected_indices)} questionnaire(s) sélectionné(s) pour l'année {year}")
        
        # Générer directement en ZIP les questionnaires sélectionnés
        result = generate_selected_questionnaires_to_zip(
            bdd_file, year, template_file, selected_indices,
            progress_callback=lambda current, total, message: update_progress(progress_bar, status_text, current, total, message)
        )
        
        if result['success']:
            st.session_state.generated_questionnaires = True
            st.session_state.generation_result = result
            st.session_state.generated_zip_data = result['zip_data']
            # Recharger la page pour afficher les résultats persistants
            st.rerun()
        else:
            st.error(f"❌ Erreur: {result['error']}")
            
    except ImportError as e:
        st.error(f"❌ Module de génération non trouvé: {e}")
        st.error("🔄 Tentative avec l'ancienne méthode...")
        # Fallback vers l'ancienne méthode
        generate_questionnaires_fallback()
    except Exception as e:
        st.error(f"❌ Erreur inattendue: {str(e)}")

def generate_questionnaires_fallback():
    """Méthode de génération de fallback"""
    
    progress_bar = st.progress(0)
    status_text = st.empty()
    
    try:
        from generateur_2025_streamlit import generate_questionnaires_for_year
        
        # Paramètres
        bdd_file = st.session_state.uploaded_bdd_file
        template_file = st.session_state.uploaded_template_file
        year = st.session_state.selected_year
        
        # Dossier de sortie temporaire
        timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        output_dir = os.path.join(
            st.session_state.temp_dir,
            f"Questionnaires_{year}_{timestamp}"
        )
        os.makedirs(output_dir, exist_ok=True)
        
        # Générer
        result = generate_questionnaires_for_year(
            bdd_file, year, template_file, output_dir,
            progress_callback=lambda current, total, message: update_progress(progress_bar, status_text, current, total, message)
        )
        
        if result['success']:
            # Créer le ZIP en mémoire
            zip_buffer = create_zip_from_directory(output_dir)
            
            st.session_state.generated_questionnaires = True
            st.session_state.generation_result = result
            st.session_state.generated_zip_data = zip_buffer
            st.rerun()
        else:
            st.error(f"❌ Erreur: {result['error']}")
            
    except Exception as e:
        st.error(f"❌ Erreur lors du fallback: {str(e)}")

def create_zip_from_directory(directory_path):
    """Créer un ZIP en mémoire à partir d'un dossier"""
    import io
    
    zip_buffer = io.BytesIO()
    
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zipf:
        if os.path.exists(directory_path):
            for root, dirs, files in os.walk(directory_path):
                for file in files:
                    if file.endswith(('.xlsx', '.xlsm')):
                        file_path = os.path.join(root, file)
                        arcname = os.path.relpath(file_path, directory_path)
                        zipf.write(file_path, arcname)
    
    return zip_buffer.getvalue()

def update_progress(progress_bar, status_text, current, total, message):
    """Mettre à jour la barre de progression"""
    progress = current / total if total > 0 else 0
    progress_bar.progress(progress)
    status_text.text(f"{message} ({current}/{total})")

if __name__ == "__main__":
    main()
